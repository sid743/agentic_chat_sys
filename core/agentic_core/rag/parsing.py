"""Turn uploaded files into text (PDF, DOCX, Markdown, plain text)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import yaml

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".xml", ".html", ".htm", ".yaml", ".yml", ".log", ".rst", ".tsv",
}


class UnsupportedDocument(ValueError):
    pass


@dataclass
class ParsedDocument:
    filename: str
    text: str
    pages: list[str] = field(default_factory=list)  # PDFs: one entry per page
    metadata: dict = field(default_factory=dict)
    is_markdown: bool = False


def split_front_matter(text: str) -> tuple[dict, str]:
    if not text.startswith("---"):
        return {}, text
    match = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.DOTALL)
    if not match:
        return {}, text
    try:
        meta = yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError:
        meta = {}
    return (meta if isinstance(meta, dict) else {}), text[match.end():]


def parse_bytes(filename: str, data: bytes) -> ParsedDocument:
    name = filename.lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""
    if ext == ".pdf" or data[:5] == b"%PDF-":
        return _parse_pdf(filename, data)
    if ext == ".docx" or (data[:2] == b"PK" and ext in ("", ".docx")):
        return _parse_docx(filename, data)
    if ext in TEXT_EXTENSIONS or not ext:
        return parse_text(filename, data.decode("utf-8", errors="replace"))
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedDocument(f"Unsupported file type: {filename}. Use PDF, DOCX, TXT or MD.") from exc
    return parse_text(filename, text)


def parse_text(filename: str, text: str) -> ParsedDocument:
    meta, body = split_front_matter(text.replace("\r\n", "\n"))
    is_md = filename.lower().endswith((".md", ".markdown")) or bool(re.search(r"(?m)^#{1,4}\s+\S", body))
    return ParsedDocument(filename=filename, text=body.strip(), metadata=meta, is_markdown=is_md)


def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        try:
            pages.append((page.extract_text() or "").strip())
        except Exception:  # noqa: BLE001 - damaged pages should not fail the upload
            pages.append("")
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        raise UnsupportedDocument(f"No extractable text in {filename} (scanned PDF?).")
    title = ""
    try:
        title = (reader.metadata.title or "") if reader.metadata else ""
    except Exception:  # noqa: BLE001
        title = ""
    return ParsedDocument(filename=filename, text=text, pages=pages, metadata={"title": title} if title else {})


def _parse_docx(filename: str, data: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append("#" * min(int(level) + 1, 4) + " " + text)
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n\n".join(lines)
    if not text.strip():
        raise UnsupportedDocument(f"No text found in {filename}.")
    title = document.core_properties.title or ""
    return ParsedDocument(filename=filename, text=text, metadata={"title": title} if title else {}, is_markdown=True)
